from io import BytesIO
from pathlib import Path

import pytest
from pypdf import PdfWriter

from dossier.ingest.errors import DecodeError, UnreadablePdfError
from dossier.ingest.parsing import parse

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"


def test_markdown_keeps_heading() -> None:
    data = (FIXTURES / "sample.md").read_bytes()
    parsed = parse("sample.md", "text/markdown", data)
    assert parsed.kind == "markdown"
    assert "Recommended dose" in parsed.pages[0].text


def test_digital_pdf_has_page_one_text() -> None:
    data = (FIXTURES / "one_page.pdf").read_bytes()
    parsed = parse("one_page.pdf", "application/pdf", data)
    assert parsed.pages[0].page_number == 1
    assert parsed.pages[0].text.strip()


def test_empty_pdf_is_unreadable(tmp_path: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    path = tmp_path / "empty.pdf"
    writer.write(path)
    with pytest.raises(UnreadablePdfError):
        parse("empty.pdf", "application/pdf", path.read_bytes())


def test_invalid_utf8_txt() -> None:
    with pytest.raises(DecodeError):
        parse("bad.txt", "text/plain", b"\xff\xfe not utf8")


def test_csv_is_one_sheet() -> None:
    parsed = parse("doses.csv", "text/csv", b"drug,dose\nWidget,10 mg\n")
    assert parsed.kind == "csv"
    assert "10 mg" in parsed.pages[0].text


def test_xlsx_sheets_are_pages() -> None:
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "Doses"
    book.active.append(["drug", "dose"])
    book.active.append(["Widget", "10 mg"])
    extra = book.create_sheet("Notes")
    extra.append(["cite the label"])
    buffer = BytesIO()
    book.save(buffer)
    parsed = parse(
        "book.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        buffer.getvalue(),
    )
    assert parsed.kind == "xlsx"
    assert len(parsed.pages) == 2
    assert "10 mg" in parsed.pages[0].text
    assert "cite the label" in parsed.pages[1].text


def test_xls_sheets_are_pages() -> None:
    parsed = parse("doses.xls", "application/vnd.ms-excel", (FIXTURES / "doses.xls").read_bytes())
    assert parsed.kind == "xlsx"
    assert len(parsed.pages) == 2
    assert "10 mg" in parsed.pages[0].text
    assert "cite the label" in parsed.pages[1].text


def test_docx_extracts_paragraphs() -> None:
    from docx import Document as WordDocument

    word = WordDocument()
    word.add_paragraph("The recommended dose is 10 mg.")
    buffer = BytesIO()
    word.save(buffer)
    parsed = parse(
        "note.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        buffer.getvalue(),
    )
    assert parsed.kind == "docx"
    assert "10 mg" in parsed.pages[0].text


def test_doc_zip_is_read_as_word() -> None:
    from docx import Document as WordDocument

    word = WordDocument()
    word.add_paragraph("The recommended dose is 10 mg.")
    buffer = BytesIO()
    word.save(buffer)
    parsed = parse("note.doc", "application/msword", buffer.getvalue())
    assert parsed.kind == "docx"
    assert "10 mg" in parsed.pages[0].text


def test_doc_ole_extracts_utf16_text() -> None:
    payload = b"hdr\x00" + "The recommended dose is 10 mg daily.".encode("utf-16le") + b"\x00tail"
    parsed = parse("legacy.doc", "application/msword", payload)
    assert parsed.kind == "docx"
    assert "10 mg" in parsed.pages[0].text
