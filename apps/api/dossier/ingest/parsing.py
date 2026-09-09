from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from io import BytesIO, StringIO

import xlrd
from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from dossier.ingest.errors import (
    DecodeError,
    UnreadableImageError,
    UnreadablePdfError,
    UnsupportedTypeError,
)
from dossier.ingest.images import image_mime

PDF_MIMES = {"application/pdf", "application/x-pdf"}
TEXT_MIMES = {"text/plain"}
MARKDOWN_MIMES = {"text/markdown", "text/x-markdown"}
CSV_MIMES = {"text/csv", "application/csv", "text/comma-separated-values"}
XLSX_MIMES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel.sheet.macroenabled.12",
}
XLS_MIMES = {
    "application/vnd.ms-excel",
    "application/vnd.ms-excel.sheet.binary.macroenabled.12",
}
DOCX_MIMES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.ms-word.document.macroenabled.12",
}
DOC_MIMES = {
    "application/msword",
    "application/vnd.ms-word",
}


@dataclass(frozen=True)
class ParsedPage:
    page_number: int
    text: str


@dataclass(frozen=True)
class ParsedDocument:
    filename: str
    mime: str
    kind: str
    pages: list[ParsedPage]


def parse(filename: str, mime: str, data: bytes) -> ParsedDocument:
    normalised = mime.split(";")[0].strip().lower()
    lower = filename.lower()
    if normalised in PDF_MIMES or lower.endswith(".pdf"):
        return _parse_pdf(filename, normalised or "application/pdf", data)
    if normalised in MARKDOWN_MIMES or lower.endswith((".md", ".markdown")):
        return _parse_text(filename, normalised or "text/markdown", data, kind="markdown")
    if normalised in TEXT_MIMES or lower.endswith(".txt"):
        return _parse_text(filename, normalised or "text/plain", data, kind="text")
    if normalised in CSV_MIMES or lower.endswith(".csv"):
        return _parse_csv(filename, normalised or "text/csv", data)
    if normalised in XLSX_MIMES or lower.endswith((".xlsx", ".xlsm")):
        return _parse_xlsx(filename, normalised or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", data)
    if normalised in XLS_MIMES or lower.endswith(".xls"):
        return _parse_xls(filename, normalised or "application/vnd.ms-excel", data)
    if normalised in DOCX_MIMES or lower.endswith((".docx", ".docm")):
        return _parse_docx(
            filename,
            normalised or "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data,
        )
    if normalised in DOC_MIMES or lower.endswith(".doc"):
        return _parse_doc(filename, normalised or "application/msword", data)
    raise UnsupportedTypeError(f"unsupported type: {mime}")


def parsed_image(filename: str, mime: str, data: bytes, text: str) -> ParsedDocument:
    resolved = image_mime(filename, mime, data)
    body = text.strip()
    if not body:
        raise UnreadableImageError("image has no extracted text")
    return ParsedDocument(
        filename=filename,
        mime=resolved,
        kind="image",
        pages=[ParsedPage(page_number=1, text=body)],
    )


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DecodeError("text files must be valid UTF-8") from exc


def _parse_text(filename: str, mime: str, data: bytes, *, kind: str) -> ParsedDocument:
    text = _decode_text(data)
    return ParsedDocument(
        filename=filename,
        mime=mime,
        kind=kind,
        pages=[ParsedPage(page_number=1, text=text)],
    )


def _parse_csv(filename: str, mime: str, data: bytes) -> ParsedDocument:
    text = _decode_text(data)
    reader = csv.reader(StringIO(text))
    rows = ["\t".join(row) for row in reader]
    return ParsedDocument(
        filename=filename,
        mime=mime,
        kind="csv",
        pages=[ParsedPage(page_number=1, text="\n".join(rows))],
    )


def _parse_xlsx(filename: str, mime: str, data: bytes) -> ParsedDocument:
    try:
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
        pages: list[ParsedPage] = []
        for index, name in enumerate(workbook.sheetnames, start=1):
            sheet = workbook[name]
            lines = [f"# {name}"]
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if cell is None else str(cell) for cell in row]
                if any(cell.strip() for cell in cells):
                    lines.append("\t".join(cells))
            pages.append(ParsedPage(page_number=index, text="\n".join(lines)))
        workbook.close()
    except Exception as exc:
        raise DecodeError("could not read spreadsheet") from exc
    if not pages:
        raise DecodeError("spreadsheet has no sheets")
    return ParsedDocument(filename=filename, mime=mime, kind="xlsx", pages=pages)


def _parse_xls(filename: str, mime: str, data: bytes) -> ParsedDocument:
    try:
        book = xlrd.open_workbook(file_contents=data)
        pages: list[ParsedPage] = []
        for index, sheet in enumerate(book.sheets(), start=1):
            lines = [f"# {sheet.name}"]
            for row_index in range(sheet.nrows):
                cells = ["" if cell is None else str(cell).strip() for cell in sheet.row_values(row_index)]
                if any(cells):
                    lines.append("\t".join(cells))
            pages.append(ParsedPage(page_number=index, text="\n".join(lines)))
    except Exception as exc:
        raise DecodeError("could not read spreadsheet") from exc
    if not pages:
        raise DecodeError("spreadsheet has no sheets")
    return ParsedDocument(filename=filename, mime=mime, kind="xlsx", pages=pages)


def _extract_ole_text(data: bytes) -> str:
    pieces = [
        match.group().decode("utf-16le").strip()
        for match in re.finditer(rb"(?:[\x20-\x7e]\x00){8,}", data)
    ]
    return "\n".join(piece for piece in pieces if " " in piece and len(piece) > 12)


def _parse_doc(filename: str, mime: str, data: bytes) -> ParsedDocument:
    if data[:2] == b"PK":
        return _parse_docx(filename, mime, data)
    text = _extract_ole_text(data)
    if not text.strip():
        raise DecodeError("could not read Word document")
    return ParsedDocument(
        filename=filename,
        mime=mime,
        kind="docx",
        pages=[ParsedPage(page_number=1, text=text)],
    )


def _parse_docx(filename: str, mime: str, data: bytes) -> ParsedDocument:
    try:
        document = WordDocument(BytesIO(data))
        parts: list[str] = [
            paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    except Exception as exc:
        raise DecodeError("could not read Word document") from exc
    if not any(part.strip() for part in parts):
        raise DecodeError("Word document has no text")
    return ParsedDocument(
        filename=filename,
        mime=mime,
        kind="docx",
        pages=[ParsedPage(page_number=1, text="\n\n".join(parts))],
    )


def _parse_pdf(filename: str, mime: str, data: bytes) -> ParsedDocument:
    reader = PdfReader(BytesIO(data))
    pages: list[ParsedPage] = []
    for index, page in enumerate(reader.pages, start=1):
        extracted = page.extract_text() or ""
        pages.append(ParsedPage(page_number=index, text=extracted))
    if not any(page.text.strip() for page in pages):
        raise UnreadablePdfError("PDF has no text layer")
    return ParsedDocument(filename=filename, mime=mime, kind="pdf", pages=pages)
